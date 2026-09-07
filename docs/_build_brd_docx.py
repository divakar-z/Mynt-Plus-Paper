#!/usr/bin/env python
"""Render the Chart Engine BRD markdown into a styled .docx (no external tools)."""
import re
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC = r"d:\Prem\Mynt-Plus-Web\docs\CHART_ENGINE_BRD.md"
OUT = r"d:\Prem\Mynt-Plus-Web\docs\CHART_ENGINE_BRD.docx"

ACCENT = RGBColor(0x2E, 0x4B, 0x8E)      # deep indigo for headings
ACCENT_LT = "DCE3F2"                       # table header fill
GREY = RGBColor(0x55, 0x55, 0x55)
BORDER = "BFBFBF"

doc = Document()

# ---- base styling -------------------------------------------------------
normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(10.5)
for sec in doc.sections:
    sec.left_margin = sec.right_margin = Inches(0.85)
    sec.top_margin = sec.bottom_margin = Inches(0.8)

def style_heading(level, size, color, before, after, bold=True):
    st = doc.styles[f"Heading {level}"]
    st.font.name = "Calibri"
    st.font.size = Pt(size)
    st.font.color.rgb = color
    st.font.bold = bold
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)

style_heading(1, 17, ACCENT, 16, 6)
style_heading(2, 13, ACCENT, 12, 4)
style_heading(3, 11.5, ACCENT, 8, 3)
style_heading(4, 11, RGBColor(0x1A, 0x1A, 0x1A), 9, 2)

def set_cell_bg(cell, hex_fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_fill)
    tcpr.append(shd)

def set_table_borders(tbl):
    tblpr = tbl._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), BORDER)
        borders.append(el)
    tblpr.append(borders)

INLINE = re.compile(r"(\*\*.+?\*\*|`.+?`)")

def add_runs(paragraph, text, base_color=None):
    """Add text honoring **bold** and `code` inline markers."""
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            r = paragraph.add_run(part[2:-2]); r.bold = True
        elif part.startswith("`") and part.endswith("`"):
            r = paragraph.add_run(part[1:-1])
            r.font.name = "Consolas"; r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor(0x33, 0x44, 0x66)
        else:
            r = paragraph.add_run(part)
        if base_color is not None:
            r.font.color.rgb = base_color

def split_row(line):
    cells = [c.strip() for c in line.strip().strip("|").split("|")]
    return cells

with open(SRC, encoding="utf-8") as f:
    lines = f.read().split("\n")

i = 0
n = len(lines)
first_h1_done = False
while i < n:
    line = lines[i]
    stripped = line.strip()

    # ---- horizontal rule -> skip (used as separators) ----
    if stripped == "---":
        i += 1
        continue

    # ---- tables ----
    if stripped.startswith("|") and i + 1 < n and re.match(r"^\|[-:\s|]+\|?$", lines[i + 1].strip()):
        header = split_row(stripped)
        i += 2  # skip header + separator
        body = []
        while i < n and lines[i].strip().startswith("|"):
            body.append(split_row(lines[i].strip()))
            i += 1
        tbl = doc.add_table(rows=1, cols=len(header))
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(tbl)
        hdr_cells = tbl.rows[0].cells
        for c, txt in zip(hdr_cells, header):
            c.text = ""
            p = c.paragraphs[0]
            add_runs(p, txt)
            for run in p.runs:
                run.bold = True
                run.font.color.rgb = ACCENT
            set_cell_bg(c, ACCENT_LT)
        for row in body:
            cells = tbl.add_row().cells
            for c, txt in zip(cells, row):
                c.text = ""
                add_runs(c.paragraphs[0], txt)
        doc.add_paragraph()
        continue

    # ---- headings ----
    if stripped.startswith("# "):
        text = stripped[2:].strip()
        if not first_h1_done:
            # Title block, centered
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(text)
            r.bold = True; r.font.size = Pt(22); r.font.color.rgb = ACCENT
            first_h1_done = True
        else:
            doc.add_heading(text, level=1)
        i += 1
        continue
    if stripped.startswith("## "):
        doc.add_heading(stripped[3:].strip(), level=2)
        i += 1
        continue
    if stripped.startswith("#### "):
        doc.add_heading(stripped[5:].strip(), level=4)
        i += 1
        continue
    if stripped.startswith("### "):
        doc.add_heading(stripped[4:].strip(), level=3)
        i += 1
        continue

    # ---- blockquote (tester notes) ----
    if stripped.startswith(">"):
        text = stripped.lstrip(">").strip()
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(6)
        add_runs(p, text, base_color=GREY)
        for run in p.runs:
            run.italic = True
        i += 1
        continue

    # ---- bullet list ----
    if re.match(r"^\s*[-*]\s+", line):
        text = re.sub(r"^\s*[-*]\s+", "", line)
        p = doc.add_paragraph(style="List Bullet")
        add_runs(p, text)
        i += 1
        continue

    # ---- numbered list ----
    if re.match(r"^\s*\d+\.\s+", line):
        text = re.sub(r"^\s*\d+\.\s+", "", line)
        p = doc.add_paragraph(style="List Number")
        add_runs(p, text)
        i += 1
        continue

    # ---- emphasised single-line metadata (e.g. **Platform:** ...) ----
    if stripped:
        p = doc.add_paragraph()
        add_runs(p, stripped)
        i += 1
        continue

    i += 1  # blank line

# ---- footer with page numbers ----
footer_p = doc.sections[0].footer.paragraphs[0]
footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = footer_p.add_run("Mynt Plus — In-House Charts BRD v1.0 · Internal Use Only · Page ")
run.font.size = Pt(8); run.font.color.rgb = GREY
fld = OxmlElement("w:fldSimple"); fld.set(qn("w:instr"), "PAGE")
footer_p._p.append(fld)

doc.save(OUT)
print("Saved:", OUT)
